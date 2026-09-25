# /// script
# requires-python = ">=3.11"
# dependencies = ["pyshp>=2.3", "openpyxl>=3.1", "pyproj>=3.6", "shapely>=2.0"]
# ///
"""
學校資料前處理：校地範圍 shapefile + 教育部名錄 -> 分縣市 GeoJSON

來源：
  各級學校範圍圖_121分帶  國土測繪中心，data.gov.tw/dataset/174606（EPSG:3826）
  115 學年各級學校名錄     教育部統計處，stats.moe.gov.tw

學級判定（見 design D3）：
  校名後綴為主要依據，教育部名錄為佐證。名錄的接合率僅八成，且
  三成以上的接合結果同名多學級（臺北市的「大安」同時有國小與國中），
  故名錄無法單獨決定學級；校名後綴則結構穩定。兩者不一致時以名錄
  為準並記入報告。

用法：
    uv run pipeline/build_schools.py
    uv run pipeline/build_schools.py --county 臺北市
"""

from __future__ import annotations

import argparse
import collections
import gzip
import hashlib
import json
import re
import sys
from pathlib import Path

import shapefile
from openpyxl import load_workbook
from pyproj import Transformer
from shapely.geometry import shape, mapping, LineString, Point
from shapely.ops import transform as shp_transform, unary_union, polygonize
from shapely.strtree import STRtree

COUNTIES = ['臺北市', '新北市', '桃園市', '臺中市', '臺南市', '高雄市', '基隆市', '新竹市',
            '嘉義市', '新竹縣', '苗栗縣', '彰化縣', '南投縣', '雲林縣', '嘉義縣', '屏東縣',
            '宜蘭縣', '花蓮縣', '臺東縣', '澎湖縣', '金門縣', '連江縣']

# 六類，對應 spec 的學校類型篩選
LEVELS = ['國小', '國中', '高中職', '大專', '特殊', '其他']

# 後綴 -> 學級。長的先比，順序即優先序。
SUFFIX = [
    ('國民中小學', 'k12'), ('國民小學', '國小'), ('國民中學', '國中'), ('中小學', 'k12'),
    ('高級中等學校', '高中職'), ('高級中學', '高中職'), ('職業學校', '高中職'), ('高級中等', '高中職'),
    ('科技大學', '大專'), ('技術學院', '大專'), ('專科學校', '大專'), ('大學', '大專'), ('學院', '大專'),
    ('特殊教育學校', '特殊'), ('啟聰學校', '特殊'), ('啟明學校', '特殊'), ('啟智學校', '特殊'),
    ('實驗中學', '高中職'), ('中學', '高中職'),
    ('國小', '國小'), ('國中', '國中'), ('高中', '高中職'),
    ('高工', '高中職'), ('高商', '高中職'), ('高職', '高中職'),
]

# 大專的分部與校區。名錄只收本校，這些接不上，但仍屬大專。
BRANCH = re.compile(r'(校區|分部|分校|推廣教育|教育中心|館)$')

MOE_FILES = [('e1_new', '國小'), ('j1_new', '國中'), ('high', '高中職'),
             ('u1_new', '大專'), ('sp1_new', '特殊')]


def fix(s: str | None) -> str:
    return (s or '').replace('台', '臺').replace('　', '').strip()


def county_of(name: str) -> str | None:
    n = fix(name)
    for c in COUNTIES:
        if n.startswith(c):
            return c
    return None


def core_of(name: str) -> tuple[str, str | None]:
    """
    取校名核心與學級。

    地圖端與名錄端的格式差異極大：
      地圖  臺北市中山區中山國民小學 / 南投縣立三光國民中學
      名錄  市立松山國小（縣市另存一欄）

    國小為「縣市＋行政區＋校名」，國中卻是「縣市＋立＋校名」——
    砍掉縣市後殘留的「立」若不處理，核心名會變成「立三光」，
    這曾使國中的接合率幾乎歸零。
    """
    n = fix(name)
    n = re.sub(r'\d+$', '', n)                         # 同校分塊的流水號
    n = re.sub(r'^.*?學校財團法人', '', n)               # 上騰學校財團法人花蓮縣上騰高工
    for c in COUNTIES:
        if n.startswith(c):
            n = n[len(c):]
            n = re.sub(r'^立', '', n)                   # 縣立／市立 的「立」
            break
    n = re.sub(r'^(國立|直轄市立|市立|縣立|私立|公立)', '', n)
    n = re.sub(r'^[^\s]{1,3}[區鄉鎮](?=\S)', '', n, count=1)   # 行政區。不含「市」以免誤砍校名
    n = re.sub(r'^(國立|直轄市立|市立|縣立|私立|公立)', '', n)
    n = re.sub(r'\(.*?\)|（.*?）', '', n)
    for suf, lvl in SUFFIX:
        if n.endswith(suf):
            return n[:-len(suf)], lvl
    return n, None


def read_moe(raw: Path):
    """讀取五份名錄。學校出現在哪一份即為哪一學級，不必解析類別欄位。"""
    rows = []
    for stem, lvl in MOE_FILES:
        path = raw / f'{stem}.xlsx'
        if not path.exists():
            print(f'  缺少名錄 {path}', file=sys.stderr)
            continue
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[0] or not row[1]:
                continue
            # 各份名錄的欄位排法不一致：多數是「[01]新北市」單格，
            # 但大專那份把代碼與縣市分成兩格（[33] / 臺北市）。
            # 故逐格比對縣市清單，而非認定某一欄。
            cty = None
            for cell in row[2:8]:
                if not isinstance(cell, str):
                    continue
                v = fix(re.sub(r'^\[\d+\]', '', cell))
                if v in COUNTIES:
                    cty = v
                    break
            rows.append((str(row[1]).strip(), cty, lvl))
        wb.close()
    return rows


def load_county_polygons(path: Path):
    """
    自 Overpass 的 admin_level=4 關係組出縣市面。

    校名判定縣市在國立學校上會失敗——「國立仁愛高級農業職業學校」
    的「仁愛」是鄉名不是縣名。座標落在哪個縣市界內才是確定的答案。

    Overpass 的 out geom 給的是組成邊界的線段，須自行組成封閉面。
    """
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding='utf-8'))
    polys, names = [], []
    for rel in data.get('elements', []):
        name = fix(rel.get('tags', {}).get('name', ''))
        if name not in COUNTIES:
            continue
        lines = [LineString([(p['lon'], p['lat']) for p in m['geometry']])
                 for m in rel.get('members', [])
                 if m.get('type') == 'way' and len(m.get('geometry', [])) > 1]
        if not lines:
            continue
        built = list(polygonize(unary_union(lines)))
        if not built:
            continue
        polys.append(unary_union(built))
        names.append(name)
    if not polys:
        return None
    return STRtree(polys), polys, names


def county_by_point(idx, x: float, y: float) -> str | None:
    """
    取包含該點且面積最小的縣市。

    臺北市完整位於新北市之內，嘉義市亦位於嘉義縣之內。自邊界線段
    組面時內孔會被填實，故一個點可能同時落在兩個縣市內；取面積較
    小者即取到飛地本身。若不這樣做，臺北市的學校會全部被判進新北市。
    """
    if idx is None:
        return None
    tree, polys, names = idx
    pt = Point(x, y)
    hits = [i for i in tree.query(pt) if polys[i].contains(pt)]
    if not hits:
        return None
    return names[min(hits, key=lambda i: polys[i].area)]


def school_id(name: str, cx: float, cy: float) -> str:
    """校地識別碼。來源的 ID 欄位非唯一（4,336 個面只有 1,523 個 ID）。"""
    h = hashlib.sha256(f'{fix(name)}\x1f{cx:.6f}\x1f{cy:.6f}'.encode('utf-8')).hexdigest()
    return h[:12]


def main() -> int:
    ap = argparse.ArgumentParser(description='學校資料前處理')
    ap.add_argument('--raw', type=Path, default=Path('data/raw'))
    ap.add_argument('--shp', type=Path,
                    default=Path('data/raw/school_shp/各級學校範圍圖_121_1150409'))
    ap.add_argument('--out', type=Path, default=Path('app/data/schools'))
    ap.add_argument('--county', help='只輸出單一縣市')
    ap.add_argument('--precision', type=int, default=6)
    ap.add_argument('--tolerance', type=float, default=0.0000027, help='簡化容差（度）')
    args = ap.parse_args()

    if not Path(str(args.shp) + '.shp').exists():
        print(f'找不到 shapefile：{args.shp}.shp', file=sys.stderr)
        return 1

    moe = read_moe(args.raw)
    by_cty, by_core = {}, {}
    for nm, cty, lvl in moe:
        c, _ = core_of(nm)
        by_cty.setdefault((cty, c), set()).add(lvl)
        by_core.setdefault(c, set()).add(lvl)
    print(f'教育部名錄 {len(moe)} 校')

    county_idx = load_county_polygons(args.raw / 'tw_counties.json')
    if county_idx:
        print(f'縣市界 {len(county_idx[2])} 個（OpenStreetMap，ODbL 1.0）')
    else:
        print('未載入縣市界，改以校名與名錄判定縣市', file=sys.stderr)

    to_wgs = Transformer.from_crs('EPSG:3826', 'EPSG:4326', always_xy=True).transform

    r = shapefile.Reader(str(args.shp), encoding='utf-8')
    stats = collections.Counter()
    disagreements, unresolved = [], []
    per_county = collections.defaultdict(list)

    for sr in r.shapeRecords():
        name = sr.record['BLOCKNAME']
        geom = shp_transform(to_wgs, shape(sr.shape.__geo_interface__))
        simple = geom.simplify(args.tolerance, preserve_topology=True)
        if simple.is_empty or simple.geom_type not in ('Polygon', 'MultiPolygon'):
            simple = geom
        c = simple.centroid

        cty = county_of(name)
        core, guess = core_of(name)
        moe_levels = by_cty.get((cty, core)) or (by_core.get(core) if cty is None else None)

        # 學級判定：名稱後綴為主，名錄佐證
        if guess and moe_levels:
            if guess in moe_levels:
                level, basis = guess, 'both'
                stats['agree'] += 1
            else:
                level, basis = sorted(moe_levels)[0], 'moe'
                stats['disagree'] += 1
                disagreements.append((name, guess, sorted(moe_levels)))
        elif guess:
            level, basis = guess, 'name'
            stats['name_only'] += 1
        elif moe_levels and len(moe_levels) == 1:
            level, basis = next(iter(moe_levels)), 'moe'
            stats['moe_only'] += 1
        elif BRANCH.search(fix(name)):
            level, basis = '大專', 'branch'      # 校區／分部，名錄只收本校
            stats['branch'] += 1
        else:
            level, basis = '其他', 'none'
            stats['unresolved'] += 1
            unresolved.append(name)

        if level == 'k12':
            level = '國中'      # 國民中小學歸國中，與 spec 六類對齊
        if level not in LEVELS:
            level = '其他'

        # 縣市判定：座標優先。校名判定在國立學校上會失敗——
        # 「國立仁愛高級農業職業學校」的「仁愛」是鄉名不是縣名。
        geo_cty = county_by_point(county_idx, c.x, c.y)
        if geo_cty:
            if cty and cty != geo_cty:
                stats['cty_conflict'] += 1
            cty = geo_cty
            stats['cty_geom'] += 1
        elif cty:
            stats['cty_name'] += 1
        else:
            cty = next((mc for mn, mc, ml in moe if core_of(mn)[0] == core and mc), None)
            if cty:
                stats['cty_moe'] += 1
        cty = cty or '未歸屬'

        per_county[cty].append({
            'type': 'Feature',
            'properties': {
                'id': school_id(name, c.x, c.y),
                'name': fix(name),
                'level': level,
                'basis': basis,
                'data_ym': sr.record['YYYYMM'],
            },
            'geometry': json.loads(json.dumps(mapping(simple)).replace('inf', '0')),
        })
        stats[level] += 1

    # 座標捨入
    def round_coords(o, nd):
        if isinstance(o, (list, tuple)):
            if o and isinstance(o[0], (int, float)):
                return [round(float(v), nd) for v in o]
            return [round_coords(i, nd) for i in o]
        return o

    args.out.mkdir(parents=True, exist_ok=True)
    index = []
    print()
    print(f"{'縣市':<8}{'校地':>6}{'MB':>8}{'gzip':>8}   各學級")
    print('-' * 74)
    for cty in sorted(per_county):
        if args.county and cty != args.county:
            continue
        feats = per_county[cty]
        for f in feats:
            f['geometry']['coordinates'] = round_coords(f['geometry']['coordinates'], args.precision)
        blob = json.dumps({'type': 'FeatureCollection', 'county': cty, 'features': feats},
                          ensure_ascii=False, separators=(',', ':')).encode('utf-8')
        (args.out / f'{cty}.geojson').write_bytes(blob)
        lv = collections.Counter(f['properties']['level'] for f in feats)
        print(f"{cty:<8}{len(feats):>6}{len(blob)/1048576:>8.2f}"
              f"{len(gzip.compress(blob,9))/1048576:>8.2f}   "
              + ' '.join(f'{k}{v}' for k, v in sorted(lv.items())))
        index.append({'county': cty, 'count': len(feats), 'bytes': len(blob),
                      'file': f'{cty}.geojson', 'levels': dict(lv)})

    (args.out / 'index.json').write_text(json.dumps(
        {'source': '國土測繪中心「各級學校範圍圖」＋教育部 115 學年學校名錄',
         'licence': '政府資料開放授權條款第 1 版',
         'counties': index}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')

    total = sum(stats[k] for k in ['agree', 'disagree', 'name_only', 'moe_only', 'branch', 'unresolved'])
    print()
    print('=== 學級判定依據 ===')
    for k, label in [('agree', '名稱與名錄一致'), ('disagree', '兩者不一致（採名錄）'),
                     ('name_only', '僅名稱可判'), ('moe_only', '僅名錄可判'),
                     ('branch', '大專分部／校區'), ('unresolved', '無法判定，歸其他')]:
        print(f'  {label:<22}{stats[k]:>6}  {100*stats[k]/total:>5.1f}%')
    print()
    print('=== 各學級校地數 ===')
    for lv in LEVELS:
        print(f'  {lv:<8}{stats[lv]:>6}')
    print()
    print('=== 縣市判定依據 ===')
    for k, label in [('cty_geom', '座標落在縣市界內'), ('cty_name', '校名前綴'),
                     ('cty_moe', '教育部名錄'), ('cty_conflict', '（座標與校名不符，採座標）')]:
        if stats[k]:
            print(f'  {label:<26}{stats[k]:>6}')

    report = {'disagreements': [{'name': n, 'by_name': g, 'by_moe': m} for n, g, m in disagreements],
              'unresolved': unresolved}
    Path('docs/school-join-report.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=1), encoding='utf-8')
    print()
    print(f'不一致 {len(disagreements)} 筆、無法判定 {len(unresolved)} 筆'
          f' -> docs/school-join-report.json')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
